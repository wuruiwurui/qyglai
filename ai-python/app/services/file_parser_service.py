"""真实文件解析服务。"""

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path

from app.core.exceptions import BusinessError
from app.schemas.ai import ParsedFileResponse


class FileParserService:
    """文件解析服务。

    支持常见办公文件格式，把文件内容统一转成原始文本，供后续合同抽取、发票解析、
    对账分析和知识库入库使用。
    """

    async def parse_upload(self, filename: str, content_type: str | None, content: bytes) -> ParsedFileResponse:
        """解析上传文件。

        Args:
            filename: 原始文件名。
            content_type: 浏览器或调用方传入的 MIME 类型。
            content: 文件二进制内容。

        Returns:
            文件解析结果。
        """

        if not content:
            raise BusinessError("上传文件为空", "EMPTY_FILE")

        extension = Path(filename).suffix.lower().lstrip(".")
        raw_text, warnings = self._parse_by_extension(extension, content)
        raw_text = raw_text.strip()
        if not raw_text:
            warnings.append("未解析出有效文本，请确认文件是否为空或是否为扫描件。")

        return ParsedFileResponse(
            filename=filename,
            contentType=content_type or "application/octet-stream",
            extension=extension,
            rawText=raw_text,
            charCount=len(raw_text),
            warnings=warnings,
        )

    def _parse_by_extension(self, extension: str, content: bytes) -> tuple[str, list[str]]:
        """根据扩展名选择解析器。"""

        if extension in {"txt", "md", "log"}:
            return self._decode_text(content), []
        if extension == "json":
            return self._parse_json(content), []
        if extension == "csv":
            return self._parse_csv(content), []
        if extension == "docx":
            return self._parse_docx(content), []
        if extension in {"xlsx", "xlsm"}:
            return self._parse_xlsx(content), []
        if extension == "pdf":
            return self._parse_pdf(content), []
        raise BusinessError(f"暂不支持的文件类型：{extension}", "UNSUPPORTED_FILE_TYPE")

    def _decode_text(self, content: bytes) -> str:
        """按常见编码解码文本文件。"""

        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    def _parse_json(self, content: bytes) -> str:
        """解析 JSON 文件。"""

        data = json.loads(self._decode_text(content))
        return json.dumps(data, ensure_ascii=False, indent=2)

    def _parse_csv(self, content: bytes) -> str:
        """解析 CSV 文件。"""

        text = self._decode_text(content)
        rows = csv.reader(StringIO(text))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

    def _parse_docx(self, content: bytes) -> str:
        """解析 Word 文档。"""

        try:
            from docx import Document
        except ImportError as exc:
            raise BusinessError("缺少 python-docx 依赖，无法解析 Word 文件", "PARSER_DEPENDENCY_MISSING") from exc

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs)

    def _parse_xlsx(self, content: bytes) -> str:
        """解析 Excel 工作簿。"""

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise BusinessError("缺少 openpyxl 依赖，无法解析 Excel 文件", "PARSER_DEPENDENCY_MISSING") from exc

        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    def _parse_pdf(self, content: bytes) -> str:
        """解析 PDF 文档。"""

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise BusinessError("缺少 pypdf 依赖，无法解析 PDF 文件", "PARSER_DEPENDENCY_MISSING") from exc

        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(page.strip() for page in pages if page.strip())


file_parser_service = FileParserService()
